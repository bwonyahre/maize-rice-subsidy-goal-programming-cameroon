"""Exploitation of the project data workbook (maize panel 2016-2024).

Outputs (../results/): t16_observed_expansion_and_yields.csv, t17_baseline_reconciliation.csv.
t17_baseline_reconciliation.csv. Provides observed_beta() used by the regional-beta scenario.
"""
import os, itertools
import openpyxl, pandas as pd

HERE = os.path.dirname(os.path.abspath(__file__))
XLSX = os.path.join(HERE, "..", "data", "projet_mais_cameroun_donnees.xlsx")
REG_EN = {"Adamaoua": "Adamawa", "Centre": "Centre", "Est": "East", "Extrême-Nord": "Far North",
          "Littoral": "Littoral", "Nord": "North", "Nord-Ouest": "Northwest", "Ouest": "West",
          "Sud": "South", "Sud-Ouest": "Southwest"}
HORIZON = 6  # years: 2024 -> 2030 planning horizon


def load_panel():
    ws = openpyxl.load_workbook(XLSX, data_only=True)["Panel_Regional_Maïs"]
    rows = []
    for r in ws.iter_rows(min_row=4, values_only=True):
        if r[0] in REG_EN and isinstance(r[2], (int, float)) and isinstance(r[3], (int, float)):
            src = str(r[5])
            if "ALT" in src or "INCOHÉRENT" in src:     # contradictory alternative series: excluded
                continue
            rows.append((REG_EN[r[0]], int(r[1]), float(r[2]), float(r[3]), src))
    df = pd.DataFrame(rows, columns=["region", "year", "area", "prod", "source"])
    # AGRISTAT is the priority source: keep it when a year has several values
    df["prio"] = df["source"].str.startswith("AGRISTAT").map({True: 0, False: 1})
    df = df.sort_values(["region", "year", "prio"]).drop_duplicates(["region", "year"])
    df["yield"] = df["prod"] / df["area"]
    return df


def observed_beta(df, horizon=HORIZON):
    """Largest relative area increase observed between two years at most `horizon` years apart."""
    out = {}
    for reg, g in df.groupby("region"):
        a = dict(zip(g.year, g.area)); best = 0.0
        for s, t in itertools.combinations(sorted(a), 2):
            if t - s <= horizon:
                best = max(best, a[t] / a[s] - 1)
        out[reg] = best
    return out


def main():
    df = load_panel()
    ob = observed_beta(df)
    rows = []
    for reg in REG_EN.values():
        g = df[df.region == reg]
        y2020 = float(g[g.year == 2020]["yield"].iloc[0]); a2020 = float(g[g.year == 2020]["area"].iloc[0])
        best = g.loc[g["yield"].idxmax()]
        rows.append({"region": reg, "years_observed": f"{g.year.min()}-{g.year.max()}",
                     "area_2020_ha": round(a2020), "yield_2020_t_ha": round(y2020, 2),
                     "best_observed_yield_t_ha": round(float(best["yield"]), 2),
                     "year_of_best_yield": int(best["year"]),
                     "max_area_expansion_within_6y": round(ob[reg], 3)})
    t9 = pd.DataFrame(rows)
    t9.to_csv(os.path.join(HERE, "..", "results", "t16_observed_expansion_and_yields.csv"), index=False)
    print(t9.to_string(index=False))

    ws = openpyxl.load_workbook(XLSX, data_only=True)["Cible_SND30_Demande"]
    series = {r[0]: r[1] for r in ws.iter_rows(values_only=True) if isinstance(r[0], int) and isinstance(r[1], (int, float)) and 2019 < r[0] < 2025}
    ag = openpyxl.load_workbook(XLSX, data_only=True)["AGRISTAT_2016-2020"]
    agr_total_2020 = [r for r in ag.iter_rows(min_row=1, max_row=20, values_only=True) if r[0] == "Total national"][0][5]
    g = (series[2024] / series[2020]) ** (1 / 4) - 1
    proj2030 = series[2024] * (1 + g) ** 6
    rec = [
        ("Maize production 2020, AGRISTAT regional total (model base)", agr_total_2020),
        ("Maize production 2020, SND30 report series", series[2020]),
        ("Maize production 2024, SND30 report series", series[2024]),
        ("Average annual growth 2020-2024 (report series)", round(g, 4)),
        ("Own extrapolation to 2030 at that rate (t)", round(proj2030)),
        ("Own extrapolation: share of the 8 Mt target", round(proj2030 / 8e6, 3)),
        ("SND30 report: projected achievement of the 2030 target", 0.341),
    ]
    pd.DataFrame(rec, columns=["item", "value"]).to_csv(
        os.path.join(HERE, "..", "results", "t17_baseline_reconciliation.csv"), index=False)
    for r in rec: print(r)


if __name__ == "__main__":
    main()
